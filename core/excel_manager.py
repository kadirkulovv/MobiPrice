import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.utils.text import slugify
from core.models import Brand, DeviceSeries, DevicePhone, DeviceIssue, IssueComparison


def export_prices_to_excel(stream_or_path=None):
    """Barcha narxlar, telefonlar va ehtiyot qismlarni formatlangan Excel jadvaliga eksport qilish"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Narxlar va Ehtiyot qismlar"

    # Styling headers
    header_fill = PatternFill(start_color="0E131F", end_color="0E131F", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="00F2FE")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    headers = [
        "ID",
        "Brend",
        "Model Seriyasi",
        "Telefon Modeli",
        "Muammo / Detal",
        "Sifat Turi (original/copy)",
        "Detal Narxi (so'm)",
        "Usta Haqqi (so'm)",
        "Umumiy Narx (so'm)",
        "Narxi (USD)",
        "Ombordagi Soni (ta)",
        "Mavjudmi (1/0)",
        "Kafolat muddati"
    ]

    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Fetch comparisons
    comparisons = IssueComparison.objects.select_related(
        'issue', 'issue__phone', 'issue__phone__series', 'issue__phone__series__brand'
    ).order_by(
        'issue__phone__series__brand__name',
        'issue__phone__series__name',
        'issue__phone__name',
        'issue__title',
        'quality_type'
    )

    for c in comparisons:
        phone = c.issue.phone
        series = phone.series
        brand = series.brand

        row = [
            c.id,
            brand.name,
            series.name,
            phone.name,
            c.issue.title,
            c.quality_type,
            int(c.part_price_uzs),
            int(c.service_fee_uzs),
            int(c.total_price_uzs),
            int(c.price_usd),
            c.quantity,
            1 if c.is_in_stock else 0,
            c.warranty or "6 oy"
        ]
        ws.append(row)
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 20

        for col_idx in range(1, len(row) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.border = thin_border
            if col_idx in [1, 6, 11, 12]:
                cell.alignment = center_align
            elif col_idx in [7, 8, 9, 10]:
                cell.alignment = right_align
                cell.number_format = '#,##0'
            else:
                cell.alignment = left_align

    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    if stream_or_path:
        wb.save(stream_or_path)
        return stream_or_path
    else:
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return out


def import_prices_from_excel(file_or_path):
    """Excel faylini o'qib bazani yangilash va yangi yozuvlarni kiritish"""
    wb = openpyxl.load_workbook(file_or_path, data_only=True)
    ws = wb.active

    created_count = 0
    updated_count = 0
    errors = []

    # Iterate rows starting from row 2
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(row):
            continue

        try:
            cid = row[0]
            brand_name = str(row[1] or '').strip()
            series_name = str(row[2] or '').strip()
            phone_name = str(row[3] or '').strip()
            issue_title = str(row[4] or '').strip()
            quality_type = str(row[5] or '').strip().lower()

            if not brand_name or not phone_name or not issue_title:
                continue

            part_price = int(float(row[6] or 0))
            service_fee = int(float(row[7] or 0))
            total_price = int(float(row[8] or (part_price + service_fee)))
            price_usd = int(float(row[9] or (total_price // 12800 if total_price else 0)))
            quantity = int(float(row[10] or 0))
            is_stock = bool(int(float(row[11] or (1 if quantity > 0 else 0))))
            warranty = str(row[12] or '6 oy').strip() if len(row) > 12 else '6 oy'

            if quality_type not in ('original', 'copy'):
                quality_type = 'original' if 'orig' in quality_type else 'copy'

            # 1. Brand
            brand_slug = slugify(brand_name) or brand_name.lower().replace(' ', '-')
            brand, _ = Brand.objects.get_or_create(
                name__iexact=brand_name,
                defaults={'name': brand_name, 'slug': brand_slug, 'is_active': True}
            )

            # 2. Series
            if not series_name:
                series_name = "Umumiy"
            series_slug = slugify(series_name) or series_name.lower().replace(' ', '-')
            series, _ = DeviceSeries.objects.get_or_create(
                brand=brand,
                name__iexact=series_name,
                defaults={'name': series_name, 'slug': series_slug}
            )

            # 3. Phone
            phone_slug = slugify(phone_name) or phone_name.lower().replace(' ', '-')
            phone, _ = DevicePhone.objects.get_or_create(
                series=series,
                name__iexact=phone_name,
                defaults={'name': phone_name, 'slug': phone_slug}
            )

            # 4. Issue
            issue, _ = DeviceIssue.objects.get_or_create(
                phone=phone,
                title__iexact=issue_title,
                defaults={'title': issue_title, 'icon_number': 1}
            )

            # 5. Comparison
            comparison = None
            if cid and str(cid).isdigit():
                comparison = IssueComparison.objects.filter(id=int(cid)).first()

            if not comparison:
                comparison = IssueComparison.objects.filter(
                    issue=issue,
                    quality_type=quality_type
                ).first()

            title_display = "Original Servis Paketi" if quality_type == 'original' else "Sifatli Dublikat (Kopiya)"

            if comparison:
                comparison.part_price_uzs = part_price
                comparison.service_fee_uzs = service_fee
                comparison.price_usd = price_usd
                comparison.quantity = quantity
                comparison.is_in_stock = is_stock
                comparison.warranty = warranty
                comparison.save()
                updated_count += 1
            else:
                IssueComparison.objects.create(
                    issue=issue,
                    quality_type=quality_type,
                    title=title_display,
                    part_price_uzs=part_price,
                    service_fee_uzs=service_fee,
                    price_usd=price_usd,
                    quantity=quantity,
                    is_in_stock=is_stock,
                    warranty=warranty,
                )
                created_count += 1


        except Exception as e:
            errors.append(f"Qator {row_idx}: {str(e)}")

    return {
        'created': created_count,
        'updated': updated_count,
        'errors': errors
    }
